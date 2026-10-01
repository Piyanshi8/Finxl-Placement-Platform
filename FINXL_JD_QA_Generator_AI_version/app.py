import json
import os

from flask import (
    Flask,
    render_template,
    request,
    jsonify,
    send_file
)

from skill_engine import extract_skills
from question_engine import generate_questions
from ai_generator import generate_ai_questions
from file_parser import extract_text_from_file

from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet

from docx import Document

from flask_cors import CORS

app = Flask(__name__)
CORS(app)


# =========================================================
# HOME
# =========================================================

@app.route("/")
def home():
    return render_template("index.html")


# =========================================================
# GENERATE
# =========================================================

@app.route("/generate", methods=["POST"])
def generate():

    try:

        print("\n===================================")
        print("        FINXL JD GENERATOR")
        print("===================================")

        mode = request.form.get(
            "mode",
            "rule"
        )

        print("Generation Mode:", mode)

        jd_text = request.form.get(
            "jd_text",
            ""
        ).strip()

        uploaded_file = request.files.get(
            "jd_file"
        )

        # -------------------------------------------------
        # FILE UPLOAD
        # -------------------------------------------------

        if uploaded_file and uploaded_file.filename:

            print(
                "Uploaded File:",
                uploaded_file.filename
            )

            jd_text = extract_text_from_file(
                uploaded_file
            ).strip()

            print(
                "Extracted Text Length:",
                len(jd_text)
            )

        # -------------------------------------------------
        # VALIDATE JD
        # -------------------------------------------------

        if not jd_text:

            return jsonify({
                "success": False,
                "error":
                    "No Job Description found. "
                    "Please paste a JD or upload a TXT, PDF, or DOCX file."
            }), 400

        print(
            "JD Length:",
            len(jd_text)
        )

        # -------------------------------------------------
        # RULE BASED
        # -------------------------------------------------

        if mode == "rule":

            print(
                "Running Rule-Based Engine..."
            )

            analysis = extract_skills(
                jd_text
            )

            questions = generate_questions(
                analysis
            )

            result = {

                "job_role":
                    analysis.get(
                        "job_role",
                        "Unknown Role"
                    ),

                "skills":
                    analysis.get(
                        "skills",
                        []
                    ),

                "total_questions":
                    len(questions),

                "questions":
                    questions
            }

        # -------------------------------------------------
        # GEMINI
        # -------------------------------------------------

        elif mode == "gemini":

            print(
                "Running Gemini AI..."
            )

            result = generate_ai_questions(
                jd_text
            )

            if "questions" in result:

                result["total_questions"] = len(
                    result["questions"]
                )

        else:

            return jsonify({
                "success": False,
                "error":
                    "Invalid generation mode."
            }), 400

        # -------------------------------------------------
        # SAVE JSON
        # -------------------------------------------------

        output_dir = "output"

        os.makedirs(
            output_dir,
            exist_ok=True
        )

        output_file = os.path.join(
            output_dir,
            "questions.json"
        )

        with open(
            output_file,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                result,
                file,
                indent=4,
                ensure_ascii=False
            )

        print(
            "JSON saved:",
            output_file
        )

        print(
            "Generated Questions:",
            result.get(
                "total_questions",
                0
            )
        )

        print(
            "===================================\n"
        )

        return jsonify({

            "success": True,

            "mode": mode,

            "data": result

        })

    except ValueError as e:

        print(
            "FILE ERROR:",
            str(e)
        )

        return jsonify({

            "success": False,

            "error": str(e)

        }), 400

    except Exception as e:

        print(
            "\n==================================="
        )

        print(
            "APPLICATION ERROR:"
        )

        print(
            str(e)
        )

        print(
            "===================================\n"
        )

        return jsonify({

            "success": False,

            "error":
                "Generation failed: " + str(e)

        }), 500


# =========================================================
# LOAD GENERATED JSON
# =========================================================

def load_generated_questions():

    output_file = os.path.join(
        "output",
        "questions.json"
    )

    if not os.path.exists(
        output_file
    ):

        return None

    with open(
        output_file,
        "r",
        encoding="utf-8"
    ) as file:

        return json.load(file)


# =========================================================
# JSON DOWNLOAD
# =========================================================

@app.route("/download/json")
def download_json():

    output_file = os.path.join(
        "output",
        "questions.json"
    )

    if not os.path.exists(
        output_file
    ):

        return jsonify({

            "success": False,

            "error":
                "Please generate interview questions first."

        }), 404

    return send_file(

        output_file,

        as_attachment=True,

        download_name=
            "FINXL_Interview_Questions.json",

        mimetype=
            "application/json"
    )


# =========================================================
# TXT DOWNLOAD
# =========================================================

@app.route("/download/txt")
def download_txt():

    result = load_generated_questions()

    if result is None:

        return jsonify({

            "success": False,

            "error":
                "Please generate interview questions first."

        }), 404


    txt_file = os.path.join(
        "output",
        "FINXL_Interview_Questions.txt"
    )


    with open(
        txt_file,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(
            "FINXL INTERVIEW PREPARATION\n"
        )

        file.write(
            "=" * 60 + "\n\n"
        )


        # Job Role

        file.write(
            "JOB ROLE\n"
        )

        file.write(
            f"{result.get('job_role', 'Unknown Role')}\n\n"
        )


        # Summary

        file.write(
            "SUMMARY\n"
        )

        file.write(
            f"Total Questions: "
            f"{result.get('total_questions', 0)}\n"
        )

        file.write(
            f"Skills Detected: "
            f"{len(result.get('skills', []))}\n\n"
        )


        # Skills

        file.write(
            "DETECTED SKILLS\n"
        )

        file.write(
            "-" * 60 + "\n"
        )


        for skill in result.get(
            "skills",
            []
        ):

            file.write(

                f"{skill.get('name', 'Unknown')} | "
                f"Priority: {skill.get('priority', 'Low')} | "
                f"Score: {skill.get('score', 0)}\n"

            )


        file.write(
            "\n\n"
        )


        # Questions

        file.write(
            "INTERVIEW QUESTIONS & ANSWERS\n"
        )

        file.write(
            "=" * 60 + "\n\n"
        )


        for index, item in enumerate(
            result.get("questions", []),
            start=1
        ):

            file.write(
                f"Q{index:02d}. "
                f"{item.get('question', '')}\n\n"
            )

            file.write(
                f"Skill: "
                f"{item.get('skill', 'General')}\n"
            )

            file.write(
                f"Difficulty: "
                f"{item.get('difficulty', 'Easy')}\n"
            )

            file.write(
                f"Type: "
                f"{item.get('type', 'Interview')}\n\n"
            )

            file.write(
                "SUGGESTED ANSWER\n"
            )

            file.write(
                f"{item.get('answer', '')}\n\n"
            )

            file.write(
                "-" * 60 + "\n\n"
            )


    return send_file(

        txt_file,

        as_attachment=True,

        download_name=
            "FINXL_Interview_Questions.txt",

        mimetype=
            "text/plain"
    )


# =========================================================
# DOCX DOWNLOAD
# =========================================================

@app.route("/download/docx")
def download_docx():

    result = load_generated_questions()

    if result is None:

        return jsonify({

            "success": False,

            "error":
                "Please generate interview questions first."

        }), 404


    doc = Document()


    # Title

    doc.add_heading(
        "FINXL Interview Preparation",
        0
    )


    # Job Role

    doc.add_paragraph(
        f"Job Role: "
        f"{result.get('job_role', 'Unknown Role')}"
    )


    doc.add_paragraph(
        f"Total Questions: "
        f"{result.get('total_questions', 0)}"
    )


    doc.add_paragraph(
        f"Skills Detected: "
        f"{len(result.get('skills', []))}"
    )


    # Skills

    doc.add_heading(
        "Detected Skills",
        level=1
    )


    for skill in result.get(
        "skills",
        []
    ):

        doc.add_paragraph(

            f"{skill.get('name', 'Unknown')} | "
            f"Priority: {skill.get('priority', 'Low')} | "
            f"Score: {skill.get('score', 0)}"

        )


    # Questions

    doc.add_heading(
        "Interview Questions & Answers",
        level=1
    )


    for index, item in enumerate(
        result.get("questions", []),
        start=1
    ):

        doc.add_heading(

            f"Q{index:02d}. "
            f"{item.get('question', '')}",

            level=2
        )


        doc.add_paragraph(

            f"Skill: "
            f"{item.get('skill', 'General')}"

        )


        doc.add_paragraph(

            f"Difficulty: "
            f"{item.get('difficulty', 'Easy')}"

        )


        doc.add_paragraph(

            f"Type: "
            f"{item.get('type', 'Interview')}"

        )


        doc.add_paragraph(
            "Suggested Answer:"
        )


        doc.add_paragraph(

            item.get(
                "answer",
                ""
            )

        )


    docx_file = os.path.join(

        "output",

        "FINXL_Interview_Questions.docx"

    )


    doc.save(
        docx_file
    )


    return send_file(

        docx_file,

        as_attachment=True,

        download_name=
            "FINXL_Interview_Questions.docx",

        mimetype=
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    )


# =========================================================
# PDF DOWNLOAD
# =========================================================

@app.route("/download/pdf")
def download_pdf():

    result = load_generated_questions()

    if result is None:

        return jsonify({

            "success": False,

            "error":
                "Please generate interview questions first."

        }), 404


    pdf_file = os.path.join(

        "output",

        "FINXL_Interview_Questions.pdf"

    )


    styles = getSampleStyleSheet()


    title_style = styles["Title"]

    heading_style = styles["Heading2"]

    normal_style = styles["BodyText"]


    doc = SimpleDocTemplate(

        pdf_file,

        pagesize=A4,

        rightMargin=45,

        leftMargin=45,

        topMargin=45,

        bottomMargin=45

    )


    elements = []


    # Title

    elements.append(

        Paragraph(

            "FINXL Interview Preparation",

            title_style

        )

    )


    elements.append(
        Spacer(1, 15)
    )


    # Job role

    elements.append(

        Paragraph(

            f"<b>Job Role:</b> "
            f"{result.get('job_role', 'Unknown Role')}",

            normal_style

        )

    )


    elements.append(

        Paragraph(

            f"<b>Total Questions:</b> "
            f"{result.get('total_questions', 0)}",

            normal_style

        )

    )


    elements.append(

        Paragraph(

            f"<b>Skills Detected:</b> "
            f"{len(result.get('skills', []))}",

            normal_style

        )

    )


    elements.append(
        Spacer(1, 20)
    )


    # Skills

    elements.append(

        Paragraph(

            "Detected Skills",

            heading_style

        )

    )


    for skill in result.get(
        "skills",
        []
    ):

        elements.append(

            Paragraph(

                f"<b>{skill.get('name', 'Unknown')}</b> — "
                f"Priority: {skill.get('priority', 'Low')} — "
                f"Score: {skill.get('score', 0)}",

                normal_style

            )

        )

        elements.append(
            Spacer(1, 5)
        )


    elements.append(
        Spacer(1, 20)
    )


    # Questions

    elements.append(

        Paragraph(

            "Interview Questions & Answers",

            heading_style

        )

    )


    for index, item in enumerate(
        result.get("questions", []),
        start=1
    ):

        elements.append(

            Paragraph(

                f"<b>Q{index:02d}. "
                f"{item.get('question', '')}</b>",

                heading_style

            )

        )


        elements.append(

            Paragraph(

                f"<b>Skill:</b> "
                f"{item.get('skill', 'General')}<br/>"

                f"<b>Difficulty:</b> "
                f"{item.get('difficulty', 'Easy')}<br/>"

                f"<b>Type:</b> "
                f"{item.get('type', 'Interview')}",

                normal_style

            )

        )


        elements.append(
            Spacer(1, 8)
        )


        elements.append(

            Paragraph(

                "<b>Suggested Answer:</b>",

                normal_style

            )

        )


        answer = (
            item.get(
                "answer",
                ""
            )
            .replace(
                "&",
                "&amp;"
            )
            .replace(
                "<",
                "&lt;"
            )
            .replace(
                ">",
                "&gt;"
            )
            .replace(
                "\n",
                "<br/>"
            )
        )


        elements.append(

            Paragraph(

                answer,

                normal_style

            )

        )


        elements.append(
            Spacer(1, 20)
        )


    doc.build(
        elements
    )


    return send_file(

        pdf_file,

        as_attachment=True,

        download_name=
            "FINXL_Interview_Questions.pdf",

        mimetype=
            "application/pdf"

    )


# =========================================================
# START SERVER
# =========================================================

if __name__ == "__main__":

    print("\n===================================")
    print("      FINXL JD Q&A GENERATOR")
    print("===================================")
    print("Server: http://127.0.0.1:5000")
    print("===================================\n")

    app.run(
        host="0.0.0.0",
        port=int(os.environ.get("PORT", 5000)),
        debug=True
    )